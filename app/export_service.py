from __future__ import annotations

from dataclasses import dataclass

from datetime import date, datetime
from pathlib import Path
import re
from typing import Iterable


from app.models import Request


_ILLEGAL_XML_CHARS_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_DDE_PATTERN = re.compile(
    r"^[@+\-]\s*("
    r"[A-Za-zА-Яа-яЁё_][A-Za-zА-Яа-яЁё0-9_.]*\s*\("
    r"|.*\|"
    r"|[A-Za-zА-Яа-яЁё0-9_]+!"
    r"|[A-Za-z]{1,3}[1-9]\d*\s*(?:[+\-*/^&=<>:]|$)"
    r"|\(\s*(?:[+\-=@A-Za-zА-Яа-яЁё_]|\d+\s*[+\-*/^&=<>])"
    r"|\d+(?:[.,]\d+)?\s*[+*/^&=<>]"
    r")",
    re.DOTALL,
)


def sanitize_excel_value(value: object) -> object:
    """Sanitize cell value to prevent CSV/Excel formula injection and XML corruption.

    - Numbers, booleans, and dates are preserved as native types.
    - None is converted to empty string.
    - Illegal XML 1.0 control characters are stripped.
    - Strings starting with '=' or '@' or dangerous formula/DDE expressions
      starting with '+', '-' are prefixed with a single quote to force literal
      interpretation in Excel. Normal phone numbers (+7...) and bullet lists
      (- ...) are preserved as clean strings.
    """
    if value is None:
        return ""
    if isinstance(value, (int, float, bool, date, datetime)):
        return value
    text = str(value)
    text = _ILLEGAL_XML_CHARS_RE.sub("", text)
    stripped = text.lstrip()
    if stripped.startswith(("=", "@")):
        return "'" + text
    if _DDE_PATTERN.search(stripped):
        return "'" + text
    return text



class ExportError(RuntimeError):
    pass



@dataclass(frozen=True, slots=True)
class ExportResult:
    path: Path
    row_count: int


class ExportService:
    HEADERS = (
        "ID",
        "Дата поступления",
        "Срок исполнения",
        "Дата исполнения",
        "Отдел",
        "Должность",
        "От кого",
        "На кого",
        "Суть запроса",
        "Примечание",
        "Статус",
        "Документ",
        "Закреплён",
    )

    def export_xlsx(self, requests: Iterable[Request], destination: str | Path) -> ExportResult:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
            from openpyxl.utils import get_column_letter
        except ImportError as error:
            raise ExportError("Для экспорта в Excel установите зависимость openpyxl.") from error

        destination = Path(destination).expanduser().resolve(strict=False)
        if destination.suffix.lower() != ".xlsx":
            destination = destination.with_suffix(".xlsx")
        destination.parent.mkdir(parents=True, exist_ok=True)

        rows = list(requests)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Реестр"
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:M{max(1, len(rows) + 1)}"

        for column, header in enumerate(self.HEADERS, 1):
            cell = sheet.cell(1, column, header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(fill_type="solid", fgColor="EAF2FF")

        for row_index, request in enumerate(rows, 2):
            values = (
                request.id,
                request.received_date,
                request.due_date,
                request.completed_date,
                request.from_department,
                request.from_position,
                request.from_person,
                request.to_person,
                request.request_subject,
                request.note,
                request.display_status(),
                request.source_document,
                "Да" if request.is_pinned else "Нет",
            )
            for column, value in enumerate(values, 1):
                sheet.cell(row_index, column, sanitize_excel_value(value))

        widths = (8, 16, 16, 16, 24, 24, 26, 26, 50, 38, 16, 48, 12)
        for column, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(column)].width = width

        temp_path = destination.with_name(f".{destination.name}.tmp.xlsx")
        try:
            workbook.save(temp_path)
            temp_path.replace(destination)
        except OSError as error:
            temp_path.unlink(missing_ok=True)
            raise ExportError(f"Не удалось сохранить Excel-файл: {error}") from error
        finally:
            workbook.close()
        return ExportResult(destination, len(rows))
