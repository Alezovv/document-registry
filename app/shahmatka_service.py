from __future__ import annotations

from pathlib import Path


def generate_shahmatka_xlsx(start_num: int, max_num: int, rows: int, destination: str | Path) -> Path:
    """Generate the printable number grid without depending on Qt.

    openpyxl is imported lazily so opening the main application does not pay
    the import-time cost unless the user actually generates a Shahmatka.
    """
    start_num = int(start_num)
    max_num = int(max_num)
    rows = int(rows)
    if start_num > max_num:
        raise ValueError("Начальное число не может быть больше конечного.")
    if rows < 1:
        raise ValueError("Количество строк должно быть положительным.")

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Side
    from openpyxl.worksheet.pagebreak import Break

    destination = Path(destination)
    total_nums = max_num - start_num + 1
    cols = (total_nums + rows - 1) // rows

    wb = Workbook()
    ws = wb.active
    ws.title = "Шахматка"
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_margins.left = 0.5
    ws.page_margins.right = 0.5
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5

    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )
    centered = Alignment(horizontal="center", vertical="center")
    for row_index in range(rows):
        for col_index in range(cols):
            value = start_num + col_index * rows + row_index
            number_cell = ws.cell(row=row_index + 1, column=col_index * 2 + 1)
            empty_cell = ws.cell(row=row_index + 1, column=col_index * 2 + 2)
            if value <= max_num:
                number_cell.value = value
            number_cell.border = thin_border
            number_cell.alignment = centered
            empty_cell.border = thin_border
            empty_cell.alignment = centered

    for col_index in range(cols):
        number_letter = ws.cell(row=1, column=col_index * 2 + 1).column_letter
        empty_letter = ws.cell(row=1, column=col_index * 2 + 2).column_letter
        ws.column_dimensions[number_letter].width = 8
        ws.column_dimensions[empty_letter].width = 12
        if (col_index + 1) % 6 == 0 and col_index != cols - 1:
            ws.col_breaks.append(Break(id=(col_index * 2 + 2)))

    try:
        wb.save(destination)
    finally:
        wb.close()
    return destination
